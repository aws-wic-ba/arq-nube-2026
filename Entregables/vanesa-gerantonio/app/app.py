import json
import os
from datetime import datetime, timezone
from io import BytesIO

from docx import Document
from flask import Flask, abort, flash, redirect, render_template, request, send_file, url_for
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import func, or_

from seed_data import ORGANIZATION, default_functions, default_mission, infer_type


db = SQLAlchemy()


class Unit(db.Model):
    __tablename__ = "units"
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(180), nullable=False, unique=True)
    unit_type = db.Column(db.String(60), nullable=False)
    parent_id = db.Column(db.Integer, db.ForeignKey("units.id"), nullable=True)
    mission = db.Column(db.Text, nullable=False, default="")
    status = db.Column(db.String(30), nullable=False, default="Vigente")
    updated_at = db.Column(db.DateTime(timezone=True), nullable=False, default=lambda: datetime.now(timezone.utc))
    parent = db.relationship("Unit", remote_side=[id], backref=db.backref("children", lazy=True, order_by="Unit.name"))
    functions = db.relationship("UnitFunction", cascade="all, delete-orphan", backref="unit", order_by="UnitFunction.sort_order")
    revisions = db.relationship("Revision", cascade="all, delete-orphan", backref="unit", order_by="Revision.version.desc()")


class UnitFunction(db.Model):
    __tablename__ = "unit_functions"
    id = db.Column(db.Integer, primary_key=True)
    unit_id = db.Column(db.Integer, db.ForeignKey("units.id"), nullable=False)
    description = db.Column(db.Text, nullable=False)
    sort_order = db.Column(db.Integer, nullable=False, default=0)


class Revision(db.Model):
    __tablename__ = "revisions"
    id = db.Column(db.Integer, primary_key=True)
    unit_id = db.Column(db.Integer, db.ForeignKey("units.id"), nullable=False)
    version = db.Column(db.Integer, nullable=False)
    snapshot = db.Column(db.Text, nullable=False)
    change_note = db.Column(db.String(300), nullable=False)
    created_at = db.Column(db.DateTime(timezone=True), nullable=False, default=lambda: datetime.now(timezone.utc))


def create_app(test_config=None):
    app = Flask(__name__)
    database_url = os.getenv("DATABASE_URL", "sqlite:///organicloud.db")
    app.config.from_mapping(
        SECRET_KEY=os.getenv("SECRET_KEY", "dev-only-change-me"),
        SQLALCHEMY_DATABASE_URI=database_url,
        SQLALCHEMY_TRACK_MODIFICATIONS=False,
    )
    if test_config:
        app.config.update(test_config)
    db.init_app(app)
    with app.app_context():
        db.create_all()
        seed_database()

    @app.get("/")
    def dashboard():
        query = request.args.get("q", "").strip()
        units_query = Unit.query
        if query:
            units_query = units_query.filter(or_(Unit.name.ilike(f"%{query}%"), Unit.mission.ilike(f"%{query}%")))
        units = units_query.order_by(Unit.name).limit(30).all() if query else []
        stats = {
            "units": Unit.query.count(),
            "functions": UnitFunction.query.count(),
            "versions": Revision.query.count(),
            "pending": Unit.query.filter_by(status="En revisión").count(),
        }
        recent = Revision.query.order_by(Revision.created_at.desc()).limit(6).all()
        roots = Unit.query.filter_by(parent_id=None).order_by(Unit.name).all()
        return render_template("dashboard.html", stats=stats, recent=recent, roots=roots, units=units, query=query)

    @app.get("/organigrama")
    def org_chart():
        roots = Unit.query.filter_by(parent_id=None).order_by(Unit.name).all()
        return render_template("org_chart.html", roots=roots)

    @app.get("/unidades/<int:unit_id>")
    def unit_detail(unit_id):
        return render_template("unit_detail.html", unit=Unit.query.get_or_404(unit_id))

    @app.route("/unidades/nueva", methods=["GET", "POST"])
    def unit_create():
        if request.method == "POST":
            unit = Unit()
            error = apply_form(unit)
            if error:
                flash(error, "error")
            else:
                db.session.add(unit)
                db.session.flush()
                replace_functions(unit)
                save_revision(unit, request.form.get("change_note") or "Creación de la unidad")
                db.session.commit()
                flash("La unidad se creó correctamente.", "success")
                return redirect(url_for("unit_detail", unit_id=unit.id))
        parents = Unit.query.order_by(Unit.name).all()
        return render_template("unit_form.html", unit=None, parents=parents)

    @app.route("/unidades/<int:unit_id>/editar", methods=["GET", "POST"])
    def unit_edit(unit_id):
        unit = Unit.query.get_or_404(unit_id)
        if request.method == "POST":
            error = apply_form(unit)
            if error:
                flash(error, "error")
            else:
                replace_functions(unit)
                unit.updated_at = datetime.now(timezone.utc)
                save_revision(unit, request.form.get("change_note") or "Actualización de misión y funciones")
                db.session.commit()
                flash("Los cambios se guardaron y se creó una nueva versión.", "success")
                return redirect(url_for("unit_detail", unit_id=unit.id))
        parents = Unit.query.filter(Unit.id != unit.id).order_by(Unit.name).all()
        return render_template("unit_form.html", unit=unit, parents=parents)

    @app.get("/unidades/<int:unit_id>/historial")
    def unit_history(unit_id):
        return render_template("history.html", unit=Unit.query.get_or_404(unit_id))

    @app.get("/unidades/<int:unit_id>/exportar.docx")
    def export_unit(unit_id):
        unit = Unit.query.get_or_404(unit_id)
        document = Document()
        document.add_heading(unit.name, 0)
        document.add_paragraph(f"Tipo: {unit.unit_type}")
        document.add_paragraph(f"Dependencia: {unit.parent.name if unit.parent else 'Máximo nivel'}")
        document.add_paragraph(f"Estado: {unit.status}")
        document.add_heading("Misión", level=1)
        document.add_paragraph(unit.mission)
        document.add_heading("Funciones", level=1)
        for item in unit.functions:
            document.add_paragraph(item.description, style="List Number")
        document.add_paragraph(f"Versión: {unit.revisions[0].version if unit.revisions else 1}")
        return docx_response(document, f"{safe_filename(unit.name)}.docx")

    @app.get("/exportar/manual.docx")
    def export_manual():
        document = Document()
        document.add_heading("Manual de Misiones y Funciones", 0)
        document.add_paragraph("Generado por OrganiCloud")
        for unit in Unit.query.order_by(Unit.name).all():
            document.add_heading(unit.name, level=1)
            document.add_paragraph(f"Tipo: {unit.unit_type} · Dependencia: {unit.parent.name if unit.parent else 'Máximo nivel'}")
            document.add_heading("Misión", level=2)
            document.add_paragraph(unit.mission)
            document.add_heading("Funciones", level=2)
            for item in unit.functions:
                document.add_paragraph(item.description, style="List Bullet")
        return docx_response(document, "manual-misiones-funciones.docx")

    @app.get("/health")
    def health():
        try:
            db.session.execute(db.select(func.count(Unit.id))).scalar()
            return {"status": "ok", "database": "ok"}
        except Exception:
            return {"status": "error", "database": "unavailable"}, 503

    def apply_form(unit):
        name = request.form.get("name", "").strip()
        mission = request.form.get("mission", "").strip()
        if not name or not mission:
            return "El nombre y la misión son obligatorios."
        duplicate = Unit.query.filter(func.lower(Unit.name) == name.lower(), Unit.id != (unit.id or 0)).first()
        if duplicate:
            return "Ya existe una unidad con ese nombre."
        parent_id = request.form.get("parent_id", type=int)
        if parent_id and unit.id and parent_id == unit.id:
            return "Una unidad no puede depender de sí misma."
        unit.name = name
        unit.unit_type = request.form.get("unit_type", "Unidad").strip()
        unit.mission = mission
        unit.status = request.form.get("status", "Vigente")
        unit.parent_id = parent_id
        return None

    def replace_functions(unit):
        unit.functions.clear()
        lines = [line.strip().lstrip("-•0123456789. ") for line in request.form.get("functions", "").splitlines()]
        for index, description in enumerate(filter(None, lines), 1):
            unit.functions.append(UnitFunction(description=description, sort_order=index))

    return app


def snapshot_for(unit):
    return json.dumps({
        "name": unit.name,
        "type": unit.unit_type,
        "parent": unit.parent.name if unit.parent else None,
        "mission": unit.mission,
        "status": unit.status,
        "functions": [item.description for item in unit.functions],
    }, ensure_ascii=False)


def save_revision(unit, note):
    version = (unit.revisions[0].version + 1) if unit.revisions else 1
    unit.revisions.append(Revision(version=version, snapshot=snapshot_for(unit), change_note=note[:300]))


def seed_database():
    if Unit.query.first():
        return

    def add_node(node, parent=None):
        if isinstance(node, str):
            name, children = node, []
        else:
            name, children = node
        unit = Unit(name=name, unit_type=infer_type(name), parent=parent, mission=default_mission(name), status="Vigente")
        db.session.add(unit)
        db.session.flush()
        for order, description in enumerate(default_functions(name), 1):
            unit.functions.append(UnitFunction(description=description, sort_order=order))
        save_revision(unit, "Versión inicial cargada desde la estructura organizativa")
        for child in children:
            add_node(child, unit)

    for root in ORGANIZATION:
        add_node(root)
    db.session.commit()


def docx_response(document, filename):
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return send_file(stream, as_attachment=True, download_name=filename, mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


def safe_filename(value):
    return "".join(char if char.isalnum() else "-" for char in value).strip("-").lower()


app = create_app()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=8000, debug=True)
